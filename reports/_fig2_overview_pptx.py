"""Figure 2 as an editable PowerPoint deck.

Replays the shape list that _fig2_overview.build() records while it draws, so
the .pptx cannot drift from the .pdf that LaTeX includes: there is one layout,
declared once, in _fig2_overview.py. Every module box, group frame, diamond,
connector and string becomes a native PowerPoint shape, so the whole diagram
can be dragged, recoloured and re-typed without touching Python.

The slide is 6.3 x 4.78in, the figure's own canvas, and all geometry is in
points measured from the slide's top-left corner, which is exactly what
_SHAPES stores.

Two deliberate differences from the PDF:
  * Fonts are Arial and Consolas rather than DejaVu Sans / DejaVu Sans Mono,
    because those two are present on the machines that open a .pptx. Line
    breaks are therefore not guaranteed to fall where the PDF puts them.
  * The white plate behind the "no" loop label is not reproduced; in the PDF
    it masks the dashed line underneath it.

Neither affects the PDF the paper compiles. This file is a re-layout master,
not a second source of truth.

Writes fig2_benchmark_overview.pptx to reports/figures/summary/.
Run with /usr/bin/python3 (python-pptx 1.0.2).
"""
import importlib.util
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).parent.parent
OUT = ROOT / "reports" / "figures" / "summary"

SANS, MONO = "Arial", "Consolas"
# a baseline sits this fraction of the font size below the line's top edge
BASELINE = 0.78


def load_shapes():
    """Run the figure's own build() and take the shape list it records."""
    spec = importlib.util.spec_from_file_location(
        "_fig2_overview", ROOT / "reports" / "_fig2_overview.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    fig = plt.figure(figsize=(mod.FIG_W, mod.FIG_H))
    ax = fig.add_axes([0, 0, 1, 1])
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")
    mod.build(ax)
    plt.close(fig)
    return mod, list(mod._SHAPES)


def rgb(hex_color):
    return RGBColor.from_string(hex_color.lstrip("#").upper())


def style_line(shape, color, lw, dash):
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(lw)
    if dash:
        ln = shape.line._get_or_add_ln()
        d = ln.makeelement(qn("a:prstDash"), {"val": "dash"})
        ln.append(d)


def add_arrow_head(shape):
    ln = shape.line._get_or_add_ln()
    ln.append(ln.makeelement(qn("a:tailEnd"),
                             {"type": "triangle", "w": "sm", "len": "med"}))


def add_rect(slide, s, rounded=True):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Pt(s["x"]), Pt(s["y"]), Pt(s["w"]), Pt(s["h"]))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(s["fill"])
    style_line(shape, s["edge"], s["lw"], s["dash"])
    if rounded:
        # roundRect's adjustment is the corner radius as a fraction of the
        # shorter side, so the same radius in points survives any box shape
        shape.adjustments[0] = min(0.5, s["r"] / min(s["w"], s["h"]))
    shape.shadow.inherit = False
    shape.text_frame.text = ""
    return shape


def add_diamond(slide, s):
    shape = slide.shapes.add_shape(MSO_SHAPE.DIAMOND, Pt(s["x"]), Pt(s["y"]),
                                   Pt(s["w"]), Pt(s["h"]))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(s["fill"])
    style_line(shape, s["edge"], s["lw"], False)
    shape.shadow.inherit = False
    return shape


def add_line(slide, s):
    pts = [(Emu(int(Pt(x))), Emu(int(Pt(y)))) for x, y in s["pts"]]
    ff = slide.shapes.build_freeform(pts[0][0], pts[0][1])
    ff.add_line_segments(pts[1:], close=False)
    shape = ff.convert_to_shape()
    shape.fill.background()
    style_line(shape, s["color"], s["lw"], s["dash"])
    shape.shadow.inherit = False
    if s["arrow"]:
        add_arrow_head(shape)
    return shape


def add_text(slide, s):
    size = s["size"]
    h = size * 1.7
    w = max(len(s["s"]) * size * 0.75, 24.0)
    if s["ha"] == "center":
        left, align = s["x"] - w / 2, PP_ALIGN.CENTER
    elif s["ha"] == "right":
        left, align = s["x"] - w, PP_ALIGN.RIGHT
    else:
        left, align = s["x"], PP_ALIGN.LEFT
    if s["va"] == "center":
        top, anchor = s["y"] - h / 2, MSO_ANCHOR.MIDDLE
    elif s["va"] == "baseline":
        top, anchor = s["y"] - BASELINE * size - (h - size * 1.2) / 2, \
            MSO_ANCHOR.MIDDLE
    else:                                            # "top"
        top, anchor = s["y"], MSO_ANCHOR.TOP

    tb = slide.shapes.add_textbox(Pt(left), Pt(top), Pt(w), Pt(h))
    tf = tb.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    para = tf.paragraphs[0]
    para.alignment = align
    run = para.add_run()
    run.text = s["s"]
    run.font.size = Pt(size)
    run.font.bold = bool(s.get("bold"))
    run.font.italic = bool(s.get("italic"))
    run.font.name = MONO if s.get("mono") else SANS
    run.font.color.rgb = rgb(s["color"])
    return tb


def main():
    mod, shapes = load_shapes()
    prs = Presentation()
    prs.slide_width = Inches(mod.FIG_W)
    prs.slide_height = Inches(mod.FIG_H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])     # blank

    counts = {}
    for s in shapes:
        kind = s["kind"]
        counts[kind] = counts.get(kind, 0) + 1
        if kind == "rect":
            add_rect(slide, s)
        elif kind == "diamond":
            add_diamond(slide, s)
        elif kind == "line":
            add_line(slide, s)
        elif kind == "text":
            add_text(slide, s)
        else:
            raise SystemExit(f"unhandled shape kind {kind!r}")

    path = OUT / "fig2_benchmark_overview.pptx"
    prs.save(path)
    print(f"wrote {path.name}  ({len(shapes)} shapes: " +
          ", ".join(f"{k} {v}" for k, v in sorted(counts.items())) +
          f"; slide {mod.FIG_W}x{mod.FIG_H}in)")


if __name__ == "__main__":
    main()
